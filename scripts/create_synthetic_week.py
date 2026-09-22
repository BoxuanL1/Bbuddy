"""Generate public, wholly synthetic fixtures; no real course data."""
import json
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1] / 'examples' / 'synthetic-week'


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    lectures = {
        'lecture-1': [
            'Today we study gradient descent. For a differentiable objective, the gradient points toward the steepest local increase. We move in the negative gradient direction.',
            'The update is x next equals x minus alpha times the gradient. Alpha is the step size. A large step can overshoot, while a small step may converge slowly.',
            'For f of x equals one half x squared, the gradient is x. Start at x equals four and use alpha equals one half: the iterates are four, two, one, and one half.',
            'This week compare two step sizes on the same quadratic and explain the difference. The assignment is due Friday; the calendar date is not stated here.'
        ],
        'lecture-2': [
            'Newton method uses the Hessian, the matrix of second derivatives. Solve H times p equals negative gradient, then update x by adding p.',
            'For the scalar quadratic one half x squared, the Hessian is one. Newton takes x to zero in one step. This special quadratic example is not a guarantee for all objectives.',
            'Computing a full Hessian and solving a linear system can be expensive. Newton is not automatically better when dimension is large.',
            'If the Hessian is singular the Newton system may not have a unique solution. Outside a suitable local region, damping or a line search may be needed.'
        ]}
    for name, texts in lectures.items():
        folder = ROOT / name
        folder.mkdir(exist_ok=True)
        body = 'WEBVTT\n\n' + '\n\n'.join(
            f'00:00:{i*5:02}.000 --> 00:00:{i*5+5:02}.000\n{text}' for i, text in enumerate(texts)) + '\n'
        (folder / 'lecture.vtt').write_text(body, encoding='utf-8')
        presentation = Presentation()
        for title, text in [
            ('Gradient descent' if name == 'lecture-1' else 'Newton method', '\n'.join(texts[:2])),
            ('Example and conditions', '\n'.join(texts[2:]))
        ]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = text
        presentation.save(folder / 'slides.pptx')
    (ROOT / 'outline.md').write_text('# SYNTH101 · Synthetic Optimization\n\nWeek 3 contains Lecture 1 (gradient descent) and Lecture 2 (Newton method).\n\nPrerequisites: derivatives and basic linear algebra. Learning goal: compare first-order updates with curvature-based updates.\n\nAll materials are synthetic test data, not an NTU lecture.', encoding='utf-8')
    scope = {'schema_version': '1.0', 'course': 'SYNTH101', 'scope': 'Week 3',
             'basis': 'The authored synthetic outline maps Week 3 to both synthetic lectures.',
             'verified': True, 'missing': [],
             'recordings': {name + '/lecture.vtt': {'verified': True,
                 'basis': 'Synthetic fixture authored with its companion slides; English text is the full fixture.',
                 'language': 'English'} for name in lectures},
             'roles': {'outline.md': 'outline'}}
    (ROOT / 'scope.json').write_text(json.dumps(scope, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
