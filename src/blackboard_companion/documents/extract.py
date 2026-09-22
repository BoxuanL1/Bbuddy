"""Optional, page-aware document extraction."""
from __future__ import annotations
from pathlib import Path

def extract_document(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires the 'documents' extra") from exc
        return [{"page": i, "text": page.extract_text() or ""} for i, page in enumerate(PdfReader(str(path)).pages, 1)]
    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("PPTX extraction requires the 'documents' extra") from exc
        result = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            result.append({"page": i, "text": "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))})
        return result
    return []
